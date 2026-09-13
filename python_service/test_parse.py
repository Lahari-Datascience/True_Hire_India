#!/usr/bin/env python3
"""
Unit test harness for isolated resume parsing.
Creates a synthetic PDF and DOCX resume, runs extraction, and asserts text content.
"""

import os
import sys
import tempfile
import pymupdf as fitz
import docx
from parse_resume import parse_resume

def test_pdf_extraction():
    print("Testing PDF extraction via PyMuPDF...")
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        temp_pdf = f.name
        
    doc = fitz.open()
    page = doc.new_page()
    test_text = """
    Aarav Sharma
    Senior Full-Stack Engineer | Bengaluru, India | aarav.sharma@example.com | +91 98765 43210
    
    Professional Summary:
    Full-stack engineer with 4+ years of experience building scalable distributed web applications using React, Next.js, Node.js, and PostgreSQL. Experienced in deploying cloud infrastructure on AWS and integrating GenAI APIs.
    
    Technical Skills:
    - Languages: TypeScript, JavaScript, Python, SQL
    - Frameworks: React, Next.js, Tailwind CSS, Express, Fastify
    - Databases & Tools: PostgreSQL, Redis, Supabase, Docker, Git, AWS (S3, Lambda)
    
    Work Experience:
    Senior Software Engineer - TechVanguard Solutions (2022 - Present)
    - Architected real-time dashboard supporting 20,000 active daily recruiters, reducing query latency by 45%.
    - Integrated LLM-based candidate screening pipeline.
    
    Education:
    B.Tech in Computer Science - RV College of Engineering, Bengaluru (2018 - 2022) | CGPA: 8.9/10
    """
    page.insert_text((50, 72), test_text, fontsize=11)
    doc.save(temp_pdf)
    doc.close()
    
    res = parse_resume(temp_pdf)
    assert res["success"] is True, f"PDF extraction failed: {res.get('error')}"
    assert "Aarav Sharma" in res["text"], "Failed to locate candidate name in extracted PDF text"
    assert "TechVanguard Solutions" in res["text"], "Failed to locate experience in extracted PDF text"
    if os.path.exists(temp_pdf):
        os.remove(temp_pdf)
    print("[OK] PDF extraction test PASSED! (Char count:", res["char_count"], ")")

def test_docx_extraction():
    print("Testing DOCX extraction via python-docx...")
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        temp_docx = f.name
        
    doc = docx.Document()
    doc.add_heading("Priya Patel", level=1)
    doc.add_paragraph("Data Scientist & Machine Learning Engineer | Hyderabad | priya.p@example.com")
    doc.add_heading("Skills", level=2)
    doc.add_paragraph("Python, PyTorch, Scikit-Learn, Pandas, NumPy, SQL, FastAPI, Docker, NLP, Gemini API")
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Machine Learning Engineer at CyberData Labs (2021 - Present)")
    doc.add_paragraph("Built recommendation systems and fine-tuned transformer models for information retrieval.")
    doc.save(temp_docx)
    
    res = parse_resume(temp_docx)
    assert res["success"] is True, f"DOCX extraction failed: {res.get('error')}"
    assert "Priya Patel" in res["text"], "Failed to locate candidate name in extracted DOCX text"
    assert "CyberData Labs" in res["text"], "Failed to locate company in extracted DOCX text"
    if os.path.exists(temp_docx):
        os.remove(temp_docx)
    print("[OK] DOCX extraction test PASSED! (Char count:", res["char_count"], ")")

if __name__ == "__main__":
    test_pdf_extraction()
    test_docx_extraction()
    print("\nAll resume extraction tests passed successfully!")
