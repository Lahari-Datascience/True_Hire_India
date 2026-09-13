#!/usr/bin/env python3
"""
TrueHireIndia - Isolated Resume Text Extraction Service
Supports: PDF (via PyMuPDF/fitz), DOCX (via python-docx), TXT
Outputs structured JSON to stdout for consumption by Next.js API route.
"""

import sys
import json
import os

def extract_from_pdf(file_path: str) -> dict:
    import pymupdf as fitz
    doc = fitz.open(file_path)
    text_chunks = []
    page_count = len(doc)
    
    for page_num in range(page_count):
        page = doc[page_num]
        text_chunks.append(page.get_text("text"))
        
    full_text = "\n".join(text_chunks).strip()
    return {
        "success": True,
        "file_type": "pdf",
        "page_count": page_count,
        "char_count": len(full_text),
        "text": full_text,
        "error": None
    }

def extract_from_docx(file_path: str) -> dict:
    import docx
    doc = docx.Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    
    # Also extract text inside tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text and cell_text not in paragraphs:
                    paragraphs.append(cell_text)
                    
    full_text = "\n".join(paragraphs).strip()
    return {
        "success": True,
        "file_type": "docx",
        "page_count": max(1, len(doc.paragraphs) // 30),
        "char_count": len(full_text),
        "text": full_text,
        "error": None
    }

def extract_from_txt(file_path: str) -> dict:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read().strip()
    return {
        "success": True,
        "file_type": "txt",
        "page_count": 1,
        "char_count": len(text),
        "text": text,
        "error": None
    }

def parse_resume(file_path: str) -> dict:
    if not os.path.exists(file_path):
        return {
            "success": False,
            "error": f"File not found: {file_path}",
            "text": ""
        }
        
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == ".pdf":
            return extract_from_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return extract_from_docx(file_path)
        elif ext in [".txt", ".rtf"]:
            return extract_from_txt(file_path)
        else:
            # Fallback attempt with PyMuPDF or plain text
            try:
                return extract_from_pdf(file_path)
            except Exception:
                return extract_from_txt(file_path)
    except Exception as e:
        return {
            "success": False,
            "error": f"Extraction error: {str(e)}",
            "file_type": ext.replace(".", ""),
            "text": ""
        }

if __name__ == "__main__":
    if len(sys.argv) < 2:
        output = {
            "success": False,
            "error": "Usage: python parse_resume.py <path_to_file>",
            "text": ""
        }
    else:
        file_path = sys.argv[1]
        output = parse_resume(file_path)
        
    print(json.dumps(output, ensure_ascii=False, indent=2))
